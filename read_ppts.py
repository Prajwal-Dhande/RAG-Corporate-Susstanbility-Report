from pptx import Presentation
import sys

def print_ppt_info(filepath):
    print(f"--- Info for {filepath} ---")
    try:
        prs = Presentation(filepath)
        for i, slide in enumerate(prs.slides):
            print(f"Slide {i+1}:")
            if slide.shapes.title:
                print(f"  Title: {slide.shapes.title.text.strip()}")
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    if shape != slide.shapes.title:
                        texts.append(shape.text.strip()[:100].replace('\n', ' '))
            if texts:
                print(f"  Texts: {texts}")
    except Exception as e:
        print(f"Error reading {filepath}: {e}")

if __name__ == "__main__":
    print_ppt_info("FINAL NEW 2st SEMINAR -1 (1).pptx")
    print_ppt_info("PPT Format for PQAI.pptx")
